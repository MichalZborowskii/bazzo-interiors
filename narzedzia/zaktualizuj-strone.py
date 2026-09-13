#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Wczytuje dane-realizacji.xlsx + zdjęcia z do-wrzucenia/ i aktualizuje podstrony realizacji.

    narzedzia/venv/bin/python narzedzia/zaktualizuj-strone.py --sprawdz   # raport, bez zmian
    narzedzia/venv/bin/python narzedzia/zaktualizuj-strone.py             # zapisuje zmiany
"""
import hashlib, html, re, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "dane-realizacji.xlsx"
WRZUC = ROOT / "do-wrzucenia"
ASSETS = ROOT / "assets"
DRY = "--sprawdz" in sys.argv

SLUGS = {
    "Mieszkanie w Olsztynie": "mieszkanie-olsztyn",
    "Kancelaria prawna": "kancelaria-prawna",
    "Salon fryzjerski": "salon-fryzjerski",
    "Mieszkanie": "mieszkanie",
}
EXTS = (".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".tif", ".tiff")
GALERIA_W, MATERIAL_W = 1600, 1200
# Zdjęcia już obecne na stronie zostają nietknięte — przepisujemy tylko to,
# co faktycznie przyszło w do-wrzucenia/ (inaczej tracimy jakość na ponownej kompresji).

problemy, zmiany = [], []

# ---------------- pomocnicze: wycinanie bloków HTML ----------------
def zakres_bloku(s, otwarcie, tag="div"):
    """Zwraca (start_tresci, koniec_tresci) dla bloku zaczynającego się `otwarcie`,
    licząc zagnieżdżone tagi — w przeciwieństwie do regexa nie urwie się na pierwszym </div>."""
    i = s.find(otwarcie)
    if i < 0:
        return None
    start = i + len(otwarcie)
    glebokosc = 1
    wzor = re.compile(rf"<{tag}\b[^>]*>|</{tag}>", re.I)
    poz = start
    while glebokosc:
        m = wzor.search(s, poz)
        if not m:
            return None
        glebokosc += 1 if not m.group(0).startswith("</") else -1
        poz = m.end()
        if glebokosc == 0:
            return start, m.start()
    return None

def podmien_blok(s, otwarcie, nowa, opis, tag="div"):
    z = zakres_bloku(s, otwarcie, tag)
    if not z:
        problemy.append(f"nie znaleziono bloku: {opis}")
        return s
    a, b = z
    return s[:a] + nowa + s[b:]

def podmien_raz(s, wzor, nowa, opis):
    nowy, ile = re.subn(wzor, lambda m: m.group(1) + nowa + m.group(3), s, count=1, flags=re.S)
    if ile != 1:
        problemy.append(f"nie udało się podmienić: {opis}")
    return nowy

# ---------------- odczyt arkusza ----------------
def komorka(ws, etykieta, kol=2):
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if c.value and str(c.value).strip() == etykieta:
            v = ws.cell(row=c.row, column=kol).value
            return str(v).strip() if v is not None else ""
    return ""

def wiersze(ws, prefix, kolumny):
    out = []
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if c.value and str(c.value).strip().startswith(prefix):
            out.append({n: (str(ws.cell(row=c.row, column=k).value).strip()
                            if ws.cell(row=c.row, column=k).value is not None else "")
                        for n, k in kolumny.items()})
    return out

def czytaj():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    dane = {}
    for tab, slug in SLUGS.items():
        if tab not in wb.sheetnames:
            problemy.append(f"brak zakładki „{tab}” w pliku Excel")
            continue
        ws = wb[tab]
        dane[slug] = {
            "nazwa": komorka(ws, "Nazwa realizacji"),
            "lokalizacja": komorka(ws, "Lokalizacja"),
            "powierzchnia": komorka(ws, "Powierzchnia"),
            "rok": komorka(ws, "Rok"),
            "zakres": komorka(ws, "Zakres"),
            "lead": komorka(ws, "Podtytuł (1 zdanie)"),
            "akapity": [p["tresc"] for p in wiersze(ws, "Akapit ", {"tresc": 2}) if p["tresc"]],
            "materialy": wiersze(ws, "Materiał ", {"element": 2, "material": 3, "kod": 4, "alt": 5}),
            "galeria": wiersze(ws, "Zdjęcie ", {"alt": 2}),
        }
    return dane

# ---------------- zdjęcia ----------------
def znajdz(folder, numer):
    if not folder.is_dir():
        return None
    for p in sorted(folder.iterdir()):
        if p.stem.strip() == str(numer) and p.suffix.lower() in EXTS:
            return p
    return None

def konwertuj(src, dst, szerokosc):
    """Konwersja na WebP z poszanowaniem orientacji EXIF (inaczej zdjęcia
    z telefonu wychodzą obrócone). HEIC z iPhone'a wymaga pillow-heif."""
    from PIL import Image, ImageOps
    if src.suffix.lower() in (".heic", ".heif"):
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
        except ImportError:
            problemy.append(f"{src.name}: brak obsługi HEIC — uruchom "
                            f"narzedzia/venv/bin/pip install pillow-heif")
            raise
    with Image.open(src) as im:
        im = ImageOps.exif_transpose(im)
        im = im.convert("RGB")
        if im.width > szerokosc:
            im = im.resize((szerokosc, round(im.height * szerokosc / im.width)), Image.LANCZOS)
        w, h = im.size
        if not DRY:
            dst.parent.mkdir(parents=True, exist_ok=True)
            im.save(dst, "WEBP", quality=82, method=6)
    return w, h

def wymiary(p):
    from PIL import Image
    with Image.open(p) as im:
        return im.size

def wersja(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()[:8]

def zdjecie(slug, rodzaj, numer, nazwa_docelowa, szerokosc):
    """Zwraca (sciezka_rel, w, h, wersja) albo None. Gdy brak nowego pliku
    w do-wrzucenia/, zostawia obecne zdjęcie ze strony."""
    podkat = "realizacje" if rodzaj == "galeria" else "materialy"
    dst = ASSETS / podkat / slug / f"{nazwa_docelowa}.webp"
    src = znajdz(WRZUC / slug / rodzaj, numer)
    if src is not None:
        w, h = konwertuj(src, dst, szerokosc)
        if DRY:
            zmiany.append(f"{src.relative_to(ROOT)} -> assets/{podkat}/{slug}/{nazwa_docelowa}.webp ({w}x{h})")
            return f"{podkat}/{slug}/{nazwa_docelowa}.webp", w, h, "nowe"
        zmiany.append(f"{src.relative_to(ROOT)} -> assets/{podkat}/{slug}/{nazwa_docelowa}.webp ({w}x{h})")
        return f"{podkat}/{slug}/{nazwa_docelowa}.webp", w, h, wersja(dst)
    if dst.exists():
        w, h = wymiary(dst)
        return f"{podkat}/{slug}/{nazwa_docelowa}.webp", w, h, wersja(dst)
    return None

# ---------------- budowa HTML ----------------
def e(s):
    return html.escape(s, quote=True)

def blok_meta(d):
    pary = [("Lokalizacja", d["lokalizacja"]), ("Powierzchnia", d["powierzchnia"]),
            ("Rok", d["rok"]), ("Zakres", d["zakres"])]
    return "".join(f'<div class="meta-item"><dt>{k}</dt><dd>{e(v) if v else "—"}</dd></div>'
                   for k, v in pary)

def nazwy_materialow(s, slug):
    """Nazwy plików faktur w kolejności, w jakiej występują na stronie."""
    return re.findall(rf'materialy/{re.escape(slug)}/([a-z0-9_-]+)\.webp', s)

def blok_materialy(s, slug, d):
    obecne = nazwy_materialow(s, slug)
    out = []
    for i, m in enumerate(d["materialy"], 1):
        if not (m["element"] or m["material"]):
            continue
        kod = m["kod"].strip().lower()
        if not re.fullmatch(r"#[0-9a-f]{6}", kod):
            problemy.append(f"{slug}: materiał {i} („{m['element']}”) ma nieprawidłowy kod koloru "
                            f"„{m['kod']}” — oczekiwano #rrggbb")
            continue
        nazwa = obecne[i-1] if i <= len(obecne) else re.sub(r"[^a-z0-9]+", "-", m["element"].lower()).strip("-") or f"material{i}"
        got = zdjecie(slug, "materialy", i, nazwa, MATERIAL_W)
        if got is None:
            problemy.append(f"{slug}: brak zdjęcia faktury dla materiału {i} "
                            f"(wrzuć do-wrzucenia/{slug}/materialy/{i}.jpg)")
            continue
        rel, w, h, v = got
        alt = m["alt"] or (f"Faktura: {m['material'].lower()}" if m["material"] else "Faktura materiału")
        out.append(
            f'<figure class="material"><img src="../assets/{rel}?v={v}" alt="{e(alt)}" loading="lazy" '
            f'width="{w}" height="{h}"><span class="swatch" style="background:{kod}"></span>'
            f'<figcaption><strong>{e(m["element"])}</strong><span>{e(m["material"])}</span>'
            f'<code>{e(kod)}</code></figcaption></figure>')
    return "".join(out), len(out)

def blok_galeria(slug, d):
    out, n = [], 0
    for i, g in enumerate(d["galeria"], 1):
        got = zdjecie(slug, "galeria", i, f"{i:02d}", GALERIA_W)
        if got is None:
            continue
        rel, w, h, v = got
        if not g["alt"]:
            problemy.append(f"{slug}: zdjęcie {i} nie ma opisu w kolumnie „Opis zdjęcia” "
                            f"— wymagane dla czytników ekranu i Google")
        n += 1
        klasa = "shot-pionowe" if h >= w else "shot-poziome"
        out.append(
            f'<figure class="shot {klasa}"><button class="photo-button" data-photo="../assets/{rel}?v={v}" '
            f'aria-label="Powiększ zdjęcie {n}"><img src="../assets/{rel}?v={v}" alt="{e(g["alt"])}" '
            f'loading="lazy" width="{w}" height="{h}"><span>↗</span></button></figure>')
    return "".join(out), n

def aktualizuj(slug, d):
    plik = ROOT / "realizacje" / f"{slug}.html"
    if not plik.exists():
        problemy.append(f"brak pliku realizacje/{slug}.html")
        return
    s = oryg = plik.read_text(encoding="utf-8")

    if d["nazwa"]:
        s = podmien_raz(s, r"(<h1>)(.*?)(</h1>)", e(d["nazwa"]), f"{slug}: tytuł h1")
        s = podmien_raz(s, r"(<title>)(.*?)( - Bazzo Interiors</title>)", e(d["nazwa"]), f"{slug}: title")
        s = podmien_raz(s, r'(<span aria-current="page">)(.*?)(</span>)', e(d["nazwa"]), f"{slug}: breadcrumb")
    if d["lead"]:
        s = podmien_raz(s, r'(<p class="lead">)(.*?)(</p>)', e(d["lead"]), f"{slug}: podtytuł")
        s = podmien_raz(s, r'(<meta name="description" content=")([^"]*)(")',
                        e(d["lead"]) + " Realizacja Bazzo Interiors.", f"{slug}: meta description")

    s = podmien_blok(s, '<dl class="project-meta">', blok_meta(d), f"{slug}: metryka", tag="dl")
    if d["akapity"]:
        s = podmien_blok(s, '<div class="project-text">',
                         "".join(f"<p>{e(p)}</p>" for p in d["akapity"]), f"{slug}: opis")

    mat, n_mat = blok_materialy(s, slug, d)
    if mat:
        s = podmien_blok(s, '<div class="material-grid">', mat, f"{slug}: materiały")
        if n_mat == 4 and all(m["kod"].strip() for m in d["materialy"][:4]):
            s = s.replace('<p class="note">Zestawienie robocze — do uzupełnienia.</p>', "")

    gal, n_gal = blok_galeria(slug, d)
    if gal:
        s = podmien_blok(s, '<div class="shots">', gal, f"{slug}: galeria")

    if s != oryg:
        if not DRY:
            plik.write_text(s, encoding="utf-8")
        zmiany.append(f"zaktualizowano realizacje/{slug}.html "
                      f"({n_mat} materiały, {n_gal} zdjęcia)")

def main():
    if not XLSX.exists():
        sys.exit(f"Nie znaleziono {XLSX.name} w katalogu projektu.")
    for slug, d in czytaj().items():
        aktualizuj(slug, d)

    print()
    if zmiany:
        print(f"ZMIANY ({len(zmiany)}):")
        for z in zmiany:
            print(f"  + {z}")
    else:
        print("Brak zmian — dane w Excelu zgodne z tym, co jest na stronie.")
    if problemy:
        print(f"\nDO UZUPEŁNIENIA ({len(problemy)}):")
        for p in problemy:
            print(f"  ! {p}")
    if DRY:
        print("\n(tryb --sprawdz — żaden plik nie został zmieniony)")
    elif zmiany:
        print("\nGotowe. Sprawdź stronę lokalnie, potem: git add -A && git commit -m \"...\" && git push")

if __name__ == "__main__":
    main()
